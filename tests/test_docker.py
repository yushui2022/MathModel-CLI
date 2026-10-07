"""Explicit opt-in, no credentials and no model calls. Tests the real built image."""
import os
import subprocess

import pytest

from mathmodel_cli.backend import mount, safety_options
from mathmodel_cli.config import CODEX_VERSION, DEFAULT_JOB, IMAGE
from mathmodel_cli.skills import ensure_skill
from mathmodel_cli.store import Store
from mathmodel_cli.workflow import Workflow

pytestmark = [pytest.mark.docker, pytest.mark.skipif(os.environ.get("MATHMODEL_TEST_DOCKER") != "1",
                                                   reason="set MATHMODEL_TEST_DOCKER=1 with the image built")]


def run_container(script, mounts=()):
    result = subprocess.run(["docker", "run", "--rm", *safety_options(DEFAULT_JOB["runtime"]),
                             "--network", "none", "--tmpfs", "/home/agent:rw,mode=1777",
                             *mounts, IMAGE, "python", "-c", script], capture_output=True, text=True,
                            encoding="utf-8", errors="replace", timeout=180)
    assert result.returncode == 0, result.stdout + result.stderr
    return result.stdout


def test_real_codex_version_and_no_host_key():
    result = run_container("import subprocess,os; print(subprocess.check_output(['codex','--version'],text=True)); "
                           "assert not os.getenv('OPENAI_API_KEY'); assert not os.getenv('CODEX_API_KEY')")
    assert CODEX_VERSION in result


def test_readonly_mount_and_nonroot(tmp_path):
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    (inputs / "read.txt").write_text("fixture")
    script = """
import os
from pathlib import Path
assert os.getuid() != 0
assert Path('/inputs/read.txt').read_text() == 'fixture'
for path in ['/inputs/forbidden.txt', '/opt/mathmodel/forbidden.txt']:
    try:
        Path(path).write_text('must fail')
    except OSError:
        pass
    else:
        raise AssertionError('write outside allowed area succeeded')
assert not Path('/var/run/docker.sock').exists()
print('isolation checks passed')
"""
    run_container(script, mount(inputs, "/inputs"))


def test_libreoffice_real_render_and_corrupt_docx():
    script = """
import subprocess
from pathlib import Path
from docx import Document
from pypdf import PdfReader
doc = Document()
doc.add_heading('Offline rendering fixture', 0)
doc.add_paragraph('This is not a competition paper. A real PDF should contain this text.')
doc.save('/tmp/fixture.docx')
result = subprocess.run(['libreoffice', '-env:UserInstallation=file:///tmp/lo-test', '--headless',
    '--convert-to', 'pdf', '--outdir', '/tmp', '/tmp/fixture.docx'], capture_output=True, timeout=120)
assert result.returncode == 0, result.stderr
reader = PdfReader('/tmp/fixture.pdf')
assert len(reader.pages) >= 1 and 'Offline' in reader.pages[0].extract_text()
Path('/tmp/corrupt.docx').write_bytes(b'broken zip')
try:
    Document('/tmp/corrupt.docx')
except Exception:
    pass
else:
    raise AssertionError('corrupt document accepted')
print('render passed')
"""
    run_container(script)


def test_real_standard_missing_evidence_is_blocked(tmp_path, job):
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    (inputs / "problem.txt").write_text("Offline test only. No evidence or paper was produced.")
    store = Store(tmp_path / "home")
    skill = ensure_skill(store.home)
    task = store.create(job, inputs, skill)
    state = Workflow(store).verify(task)
    assert state["status"] == "BLOCKED"
    assert state["verification"] is None
    with pytest.raises(ValueError, match="accepted snapshot"):
        Workflow(store).export(task, tmp_path / "must-not-export")
